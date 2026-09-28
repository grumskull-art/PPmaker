from copy import deepcopy
from pathlib import Path
import pytest
from pydantic import ValidationError
from pptx import Presentation
from powerpoint_app.domain.models import SlidePlan
from powerpoint_app.planning.lookup import index_sections, bind_lookup_navigation
from powerpoint_app.rendering import PptxRenderer
from powerpoint_app.quality.navigation import audit_navigation


def catalog():
    return {'navigation_groups':[{'id':'current','label':'Strøm','description':'I [A]'}],
        'quantities':{'I':'I [A]','U':'U [V]','R':'R [Ω]'},'situations':{'dc':'Ohmsk DC'},
        'entries':[{'id':'I01','seek':'I [A]','given':'U [V], R [Ω]','condition':'Ohmsk DC',
        'latex':r'I=\frac{U}{R}','lookup':{'group':'current','seek_key':'I','given_sets':[['U','R']],'situations':['dc']}}]}


def slides(db):
    e=db['entries'][0]
    return [dict(id='start',layout='lookup_start',title='Start',elements=[]),
        dict(id='index-current-1',layout='lookup_table',title='Strøm',elements=[dict(id='idx',type='table',headers=['Søger','Har opgivet','Situation / betingelse','Formel','Side'],rows=[['I01 · I [A]','U, R','Ohmsk DC',e['latex'],'—']])]),
        dict(id='card',layout='lookup_cards',title='Strøm',elements=[dict(id='I01-header',type='text',text='I01 · I [A]\nU [V], R [Ω]\nOhmsk DC'),dict(id='I01-formula',type='formula',latex=e['latex']),dict(id='I01-details',type='text',text='24 V / 12 Ω = 2 A')])]


def plan(s):
    return SlidePlan.model_validate(dict(schema_version='1.1',deck=dict(title='Test',audience='BM4',duration_minutes=1,page_format='a4_landscape'),slides=s))


def test_navigation_survives_inserted_cover_and_reordered_destinations(tmp_path):
    db=catalog();s=slides(db)
    s.insert(0,dict(id='cover',layout='lookup_table',title='Forside',elements=[dict(id='cover-table',type='table',headers=['Bemærk'],rows=[['Introduktion']])]))
    s[2],s[3]=s[3],s[2]
    positions=bind_lookup_navigation(s,db)
    assert db['entries'][0]['page']==positions['card']==3
    output=tmp_path/'nav.pptx';PptxRenderer(tmp_path).render(plan(s),output)
    deck=Presentation(output);index=deck.slides[positions['index-current-1']-1]
    link=next(sh for sh in index.shapes if sh.name=='nav-entry-I01')
    assert link.click_action.target_slide.slide_id==deck.slides[positions['card']-1].slide_id
    assert audit_navigation(plan(s),db,output)['index_entries']==1
    link.click_action.target_slide=deck.slides[0];deck.save(output)
    with pytest.raises(ValueError,match='Forkert PPTX-mål'):
        audit_navigation(plan(s),db,output)


def test_missing_duplicate_and_dangling_index_targets_rejected():
    db=catalog();s=slides(db);s[-1]['elements'][1]['id']='other-formula'
    with pytest.raises(ValueError):bind_lookup_navigation(s,db)
    db=catalog();s=slides(db);s[1]['elements'][0]['rows']*=2
    with pytest.raises(ValueError):bind_lookup_navigation(s,db)
    db=catalog();s=slides(db);bind_lookup_navigation(s,db)
    s[1]['navigation'][0]['target_slide_id']='missing'
    with pytest.raises(ValidationError):plan(s)


def test_new_formula_is_indexed_without_hardcoded_id_order():
    db=catalog();new=deepcopy(db['entries'][0]);new['id']='I99';db['entries'].append(new)
    assert [e['id'] for e in index_sections(db)[0]['entries']]==['I01','I99']
    db['entries'][1]['lookup']['situations']=['unknown']
    with pytest.raises(ValueError):index_sections(db)
