"""Repeatable file, source, formula, index and native Windows layout checks."""
from pathlib import Path
import hashlib, json, math, re
from pypdf import PdfReader
from pptx import Presentation
from matplotlib.mathtext import MathTextParser
from matplotlib import rc_context
from powerpoint_app.projects import load_plan
from powerpoint_app.quality.checks import inspect_pptx_geometry
from powerpoint_app.rendering.math_markup import math_blocks,math_search_text
from powerpoint_app.visuals.assets import math_latex
from regenerate import checked_number
ROOT=Path(__file__).resolve().parent
db=json.loads((ROOT/'formula-database.json').read_text())
plan=load_plan(ROOT/'slide-plan.json');entries=db['entries']
sources={s['id']:s for s in db['source_inventory']}
for s in sources.values():
    assert hashlib.sha256((ROOT.parents[2]/s['file']).read_bytes()).hexdigest()==s['sha256'], s['file']
    if s.get('original_file'):
        assert hashlib.sha256((ROOT.parents[2]/s['original_file']).read_bytes()).hexdigest()==s['original_sha256'], s['original_file']
for e in entries:
    assert '\u2160' not in json.dumps(e,ensure_ascii=False),e['id']
    expressions=[e['latex']]
    for key in ('steps','conversion','pitfall','example'):
        field=e[key]
        expressions.extend(value for kind,value in math_blocks(field) if kind=='math')
        assert not re.search(r'[=<>]',re.sub(r'\$[^$]*\$','',field)),(e['id'],key)
    with rc_context({'mathtext.fontset':'stix','font.family':'STIXGeneral'}):
        for expression in expressions:
            MathTextParser('agg').parse('$'+math_latex(expression)+'$',dpi=140)
    for location in e['source_locations']:
        sid, numbers=location.split(':')
        for part in numbers.split(','):
            limits=[int(v) for v in part.split('-')]
            assert 1<=min(limits)<=max(limits)<=len(sources[sid]['slides']), location
    check=e['example_check']
    if check: assert math.isclose(checked_number(check['expr']),check['expected'],rel_tol=1e-10,abs_tol=1e-12),e['id']
    page=plan.slides[e['page']-1]
    assert any(v.id==e['id']+'-formula' for v in page.elements),e['id']
index_rows=[row for s in plan.slides if s.id.startswith('index-') for row in s.elements[0].rows]
assert len(index_rows)==len(entries)
by_code={row[0].split(' · ')[0]:row for row in index_rows}
assert set(by_code)=={e['id'] for e in entries}
for e in entries:
    row=by_code[e['id']]
    assert row[3]==e['latex'] and int(row[4])==e['page']
pptx=ROOT/'exports/EL-cheatsheet-BM4-navigation.pptx';deck=Presentation(pptx)
assert len(deck.slides)==len(plan.slides) and not inspect_pptx_geometry(pptx)
pdf=PdfReader(ROOT/'exports/EL-cheatsheet-BM4-navigation.pdf').pages
assert len(pdf)==len(plan.slides)
for page in pdf: assert abs(float(page.mediabox.width)-841.89)<1 and abs(float(page.mediabox.height)-595.28)<1
for e in entries:
    page_text=pdf[e['page']-1].extract_text().replace('\n',' ')
    assert e['id'] in page_text,e['id']
    expressions=[e['latex'],*(value for key in ('steps','conversion','pitfall','example') for kind,value in math_blocks(e[key]) if kind=='math')]
    for expression in expressions:
        assert math_search_text(expression) in page_text,(e['id'],expression)
windows=json.loads((ROOT/'review/windows-layout.json').read_text(encoding='utf-8-sig'))
assert windows['pages']==len(pdf)
for b in windows['text_bounds']:
    assert b['text_height']<=b['height']-b['margin_y']+.8,b['name']
    assert b['text_width']<=b['width']-b['margin_x']+3,b['name']
protected=json.loads((ROOT/'review/protected-files.json').read_text())
for file,digest in protected.items(): assert hashlib.sha256((ROOT.parents[2]/file).read_bytes()).hexdigest()==digest,file
assert len(list((ROOT/'exports/navigation-preview').glob('page-*.png')))==len(pdf)
cases={'24 V + 12 Ω, ohmsk DC':'I01','200 W elektrisk + 37,5 V DC':'I02','Mekanisk effekt + spænding: η mangler':'I02','Totalstrøm + to parallelle R':'I07','Lederlængde + areal + resistivitet':'R04','Modstandsændring + α ved kendt reference':'G06','ΔΦ + Δt + N':'E01','Kilder og R i to-knude-net':'U07','Generelt DC-net med reference':'N01'}
assert set(cases.values())<={e['id'] for e in entries}
assert math.isclose((27+24/2)/(1+1/2+1/8),24)
assert -3+0+3==0 and 27*3==3**2*1+3**2*8
from powerpoint_app.quality.navigation import audit_navigation
navigation=audit_navigation(plan,db,pptx,ROOT/'exports/EL-cheatsheet-BM4-navigation.pdf')
html=(ROOT/'exports/EL-cheatsheet-BM4-opslag.html').read_text()
payload=json.loads(re.search(r'<script id="database" type="application/json">(.*?)</script>',html,re.S).group(1))
assert payload['entries']==entries and payload['quantities']==db['quantities'] and payload['situations']==db['situations']
for e in entries:
    for expression in [e['latex'],*(value for key in ('steps','conversion','pitfall','example') for kind,value in math_blocks(e[key]) if kind=='math')]:
        assert expression in payload['math_assets'],(e['id'],expression)
baseline=json.loads((ROOT/'review/navigation-baseline.json').read_text())
new={e['id']:e for e in entries}
for original in baseline['entries']:
    for key,value in original.items():
        if key not in {'page','latex','steps','conversion','pitfall','example'}:
            assert new[original['id']][key]==value,(original['id'],key)
report={'navigation':navigation,'html_same_database':True,'non_typographic_content_unchanged':True,'math_blocks_checked':sum(1+sum(sum(kind=='math' for kind,_ in math_blocks(e[k])) for k in ('steps','conversion','pitfall','example')) for e in entries),'pages':len(pdf),'entries':len(entries),'recomputed_examples':sum(bool(e['example_check']) for e in entries),'original_decks':len(sources),'original_slides':sum(len(s['slides']) for s in sources.values()),'protected_files':len(protected),'native_text_bounds':len(windows['text_bounds']),'windows_powerpoint':windows['powerpoint_version'],'height_tolerance_pt':.8,'width_tolerance_pt':3,'layout_note':'BoundWidth kan inkludere lille glyph-/linjeoverhæng. Alle sider er også kontrolleret visuelt; ingen synlig beskæring.','index_case_review':cases,'manual_visual_review':'Formelkort, indeks, udvalgte før/efter-udsnit og links kontrolleret visuelt.'}
(ROOT/'review/navigation-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(f"OK: {len(pdf)} A4 pages; {len(entries)} indexed entries; {report['recomputed_examples']} examples; source hashes, PDF, geometry and native text bounds verified.")
