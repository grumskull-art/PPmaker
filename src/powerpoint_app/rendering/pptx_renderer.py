from __future__ import annotations

from pathlib import Path
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from powerpoint_app.domain.models import (
    ChartElement, CircuitElement, FormulaElement, ImageElement, ProcessElement, SlidePlan,
    TableElement, TextElement, ThreeSourceCircuitElement, WarningElement,
)
from powerpoint_app.rendering.theme import Theme, load_theme
from powerpoint_app.visuals.assets import chart_png, formula_png
from powerpoint_app.rendering.teaching import presentation_frames


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def default_template_path() -> Path | None:
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[3]))
    candidate = base / "templates" / "martec_inspired.pptx"
    return candidate if candidate.is_file() else None


class PptxRenderer:
    def __init__(self, project_root: Path, theme: Theme | None = None, template_path: Path | None = None):
        self.root = project_root
        self.theme = theme or load_theme()
        project_template = project_root / "template.pptx"
        self.template_path = template_path or (project_template if project_template.is_file() else default_template_path())
        self.prs = Presentation(str(self.template_path)) if self.template_path else Presentation()
        self.prs.slide_width = Inches(13.333333)
        self.prs.slide_height = Inches(7.5)

    def render(self, plan: SlidePlan, output: Path, cancel_check=None, mode="auto") -> Path:
        problems = plan.check_assets(self.root)
        if problems:
            raise ValueError("; ".join(problems))
        frames = presentation_frames(plan, mode)
        for index, frame in enumerate(frames, 1):
            if cancel_check and cancel_check():
                raise RuntimeError("Eksporten blev annulleret.")
            self._slide(frame.slide, index, len(frames), plan, frame)
        output.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(output)
        return output

    def _slide(self, spec, number: int, total: int, plan: SlidePlan, frame):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid(); bg.fore_color.rgb = _rgb(self.theme.background)
        self._rect(slide, 0, 0, .18, 7.5, self.theme.accent)
        self._text(slide, spec.title, .65, .35, 11.8, .55, 32, self.theme.navy, bold=True)
        self._text(slide, f"{number}/{total}", 11.95, 7.05, .7, .2, 10, "687887", align=PP_ALIGN.RIGHT)
        self._text(slide, plan.deck.audience, .65, 7.05, 6, .2, 10, "687887")
        used_sources = sorted({source for element in spec.elements for source in element.source_ids})
        if used_sources:
            self._text(slide, "Kilder: " + ", ".join(used_sources), 6.9, 7.05, 4.7, .2, 9, "687887", align=PP_ALIGN.RIGHT)
        visible_formulas = [e for e in spec.elements if isinstance(e, FormulaElement) and e.id not in frame.hidden_ids]
        current_formula = visible_formulas[-1] if visible_formulas else None
        self._focus = (current_formula.diagram_id, current_formula.branch_index) if current_formula else (None, None)
        question = spec.teaching.question if spec.teaching else None
        if question:
            self._panel(slide, spec.elements, .65, 1.3, 5.4, 5.3)
            self._text(slide, question.prompt, 6.5, 1.3, 6.0, 1.65, 26, self.theme.navy, bold=True)
            if frame.show_answer:
                self._text(slide, question.answer, 6.5, 3.0, 6.0, 1.35, 26, self.theme.text)
                self._text(slide, question.explanation, 6.5, 4.45, 6.0, 2.2, 22, self.theme.text)
            else:
                if question.options:
                    choices = "\n".join(f"{chr(65 + i)}. {option}" for i, option in enumerate(question.options))
                    self._text(slide, choices, 6.5, 2.95, 6.0, 2.5, 20, self.theme.text)
                self._text(slide, f"Tænk selv / drøft med sidemanden: {question.wait_seconds} sekunder", 6.5, 5.5, 6.0, 1.0, 22, self.theme.text)
        else:
            self._render_layout(slide, spec)
        for shape in list(slide.shapes):
            if shape.name in frame.hidden_ids:
                shape._element.getparent().remove(shape._element)
        notes = slide.notes_slide.notes_text_frame
        if notes is not None:
            notes.text = spec.speaker_notes
            if spec.teaching and spec.teaching.assumptions:
                notes.text += "\nForudsætninger: " + "; ".join(spec.teaching.assumptions)
            if question:
                notes.text += f"\nSpørgsmål: {question.prompt}\nVent: {question.wait_seconds} sekunder.\nSvar: {question.answer}\nForklaring: {question.explanation}"
                if question.options:
                    notes.text += "\nValgmuligheder: " + "; ".join(question.options)
                    notes.text += f"\nKorrekt valg: {chr(65 + question.correct_option)}."
                if question.discussion_prompt:
                    notes.text += "\nSamtale efter individuelt svar: " + question.discussion_prompt

    def _render_layout(self, slide, spec):
        elements = spec.elements
        if spec.layout == "title":
            self._text(slide, spec.objective or "Undervisningspræsentation", 1.2, 2.5, 10.9, 1.5, 30, self.theme.text, align=PP_ALIGN.CENTER)
            self._rect(slide, 4.3, 4.3, 4.7, .08, self.theme.accent)
        elif spec.layout == "key_figure":
            visual = next((e for e in elements if isinstance(e, (ProcessElement, ImageElement, ChartElement, CircuitElement, ThreeSourceCircuitElement))), None)
            if isinstance(visual, ProcessElement): self._process(slide, visual)
            elif visual is not None: self._element(slide, visual, 1.2, 1.35, 10.9, 3.65)
            self._panel(slide, [e for e in elements if e is not visual], 1.2, 5.05, 10.9, 1.35)
        elif spec.layout == "two_columns" or spec.layout == "comparison":
            self._panel(slide, elements[: max(1, (len(elements)+1)//2)], .65, 1.25, 5.85, 5.45)
            self._panel(slide, elements[max(1, (len(elements)+1)//2):], 6.82, 1.25, 5.85, 5.45)
        elif spec.layout == "process_steps":
            process = next((e for e in elements if isinstance(e, ProcessElement)), None)
            if process:
                self._process(slide, process)
            self._panel(slide, [e for e in elements if e is not process], .8, 5.35, 11.7, 1.25)
        elif spec.layout == "formula_steps":
            formulas = [e for e in elements if isinstance(e, FormulaElement)]
            wide_circuit = any(isinstance(e, ThreeSourceCircuitElement) for e in elements)
            self._panel(slide, [e for e in elements if not isinstance(e, FormulaElement)], .7, 1.25, 5.45 if wide_circuit else 4.3, 5.35)
            self._panel(slide, formulas, 6.4 if wide_circuit else 5.25, 1.25, 6.25 if wide_circuit else 7.4, 5.35)
        elif spec.layout == "chart_table":
            self._panel(slide, elements, .65, 1.15, 12.0, 5.75)
        elif spec.layout in {"agenda", "summary"}:
            self._panel(slide, elements, 1.25, 1.25, 10.8, 5.45, numbered=True)
        else:
            self._panel(slide, elements, .65, 1.15, 12.0, 5.75)

    def _panel(self, slide, elements, x, y, w, h, numbered=False):
        if not elements:
            return
        gap = .14
        each = max(.7, (h - gap * (len(elements)-1)) / len(elements))
        cursor = y
        for i, element in enumerate(elements):
            self._element(slide, element, x, cursor, w, each, f"{i+1}. " if numbered else "")
            cursor += each + gap

    def _element(self, slide, element, x, y, w, h, prefix=""):
        if isinstance(element, TextElement):
            self._text(slide, prefix + element.text, x, y, w, h, 24, self.theme.text, name=element.id)
        elif isinstance(element, WarningElement):
            before = len(slide.shapes)
            self._rect(slide, x, y, w, h, self.theme.warning, radius=True, name=element.id)
            self._text(slide, prefix + element.text, x+.2, y+.12, w-.4, h-.24, 21, self.theme.text, bold=True)
            slide.shapes.add_group_shape(list(slide.shapes)[before:]).name = element.id
        elif isinstance(element, CircuitElement):
            self._circuit(slide, element, x, y, w, h)
        elif isinstance(element, ThreeSourceCircuitElement):
            self._three_source_circuit(slide, element, x, y, w, h)
        elif isinstance(element, FormulaElement):
            image = formula_png(element, self.root / "cache", self.theme)
            shape = self._picture(slide, image, x, y, w, h)
            shape.name = element.id
            self._set_alt_text(shape, element.latex)
        elif isinstance(element, ImageElement):
            shape = self._picture(slide, self.root / element.asset, x, y, w, h)
            shape.name = element.id
            self._set_alt_text(shape, element.alt_text)
        elif isinstance(element, TableElement):
            self._table(slide, element, x, y, w, h)
        elif isinstance(element, ChartElement):
            image = chart_png(element, self.root / "cache", self.theme)
            shape = self._picture(slide, image, x, y, w, h)
            shape.name = element.id
        elif isinstance(element, ProcessElement):
            self._process(slide, element)

    def _picture(self, slide, path, x, y, w, h):
        from PIL import Image
        with Image.open(path) as image:
            ratio = image.width / image.height
        pw, ph = min(w, h * ratio), min(h, w / ratio)
        return slide.shapes.add_picture(str(path), Inches(x+(w-pw)/2), Inches(y+(h-ph)/2), width=Inches(pw), height=Inches(ph))

    def _circuit(self, slide, element, x, y, w, h):
        from pptx.enum.shapes import MSO_CONNECTOR
        before = len(slide.shapes)
        if h < 2.4 or w < (4.0 if len(element.resistances) == 2 else 5.2):
            raise ValueError("Kredsløbet kræver mere plads. Brug én figur i en kolonne eller key_figure.")
        top, bottom = y + max(.7, h*.16), y + h*.83
        left, right = x + w*.18, x + w*.88
        def line(x1, y1, x2, y2, color=None):
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
            shape.line.color.rgb = _rgb(color or self.theme.navy)
            shape.line.width = Pt(3 if color else 2)
        line(left, top, right, top); line(left, bottom, right, bottom)
        radius = min(w*.075, h*.13)
        mid = (top+bottom)/2
        line(left, top, left, mid-radius); line(left, mid+radius, left, bottom)
        supply = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left-radius), Inches(mid-radius), Inches(2*radius), Inches(2*radius))
        supply.fill.solid(); supply.fill.fore_color.rgb = _rgb(self.theme.background)
        supply.line.color.rgb = _rgb(self.theme.navy)
        self._text(slide, "+\n−", left-radius, mid-radius, radius*2, radius*2, 20, self.theme.navy, align=PP_ALIGN.CENTER)
        self._text(slide, f"U = {element.voltage:g} V", x, bottom+.04, w*.43, .45, 20, self.theme.text)
        self._text(slide, "A (+)", left-.15, top-.5, 1.2, .45, 20, self.theme.navy)
        self._text(slide, "B (0 V)", right-.3, bottom+.04, 1.3, .45, 20, self.theme.navy)
        for i, resistance in enumerate(element.resistances):
            bx = left + (right-left)*(i+1)/len(element.resistances)
            rh = min(h*.28, 1.0); rw = min(w*.055, .35)
            active = self._focus == (element.id, i+1)
            color = self.theme.accent if active else self.theme.navy
            line(bx, top, bx, mid-rh/2, color if active else None); line(bx, mid+rh/2, bx, bottom, color if active else None)
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(bx-rw/2), Inches(mid-rh/2), Inches(rw), Inches(rh))
            shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(self.theme.pale)
            shape.line.color.rgb = _rgb(color)
            shape.line.width = Pt(3 if active else 1)
            label_width = min(1.1, (right-left)/len(element.resistances)-.08)
            self._text(slide, f"R{i+1}\n{resistance:g} Ω", bx-label_width/2, top-.65, label_width, .60, 18 if len(element.resistances) <= 2 else 15, color, bold=active, align=PP_ALIGN.CENTER)
        group = slide.shapes.add_group_shape(list(slide.shapes)[before:])
        group.name = element.id
        self._set_alt_text(group, f"Ideel {element.voltage:g} V DC-kilde med parallelle modstande: {element.resistances} ohm. Fælles knuder A og B.")

    def _three_source_circuit(self, slide, element, x, y, w, h):
        """Draw the bounded exam topology as editable PowerPoint shapes."""
        from pptx.enum.shapes import MSO_CONNECTOR
        if w < 5 or h < 2.8:
            raise ValueError("Kredsløb med tre kilder kræver mindst 5 × 2,8 tommer.")
        before = len(slide.shapes)
        left, a, c, d, b = [x+w*f for f in (.08, .29, .50, .72, .94)]
        top, mid, bottom = [y+h*f for f in (.20, .51, .78)]
        font = 13 if w < 7 else 16

        def line(x1, y1, x2, y2, accent=False):
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
            shape.line.color.rgb = _rgb(self.theme.accent if accent else self.theme.navy)
            shape.line.width = Pt(3 if accent else 2)

        def resistor(x1, x2, yy, n):
            center = (x1+x2)/2
            width = min(.52, (x2-x1)*.40)
            height = .22
            active = self._focus == (element.id, n)
            line(x1, yy, center-width/2, yy, active)
            line(center+width/2, yy, x2, yy, active)
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(center-width/2), Inches(yy-height/2), Inches(width), Inches(height))
            shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(self.theme.pale)
            shape.line.color.rgb = _rgb(self.theme.accent if active else self.theme.navy)
            shape.line.width = Pt(3 if active else 1)
            label_y = yy - (.48 if yy == top else .44)
            self._text(slide, f"R{n}  {element.resistances[n-1]:g} Ω", center-.64, label_y, 1.28, .38, font, self.theme.accent if active else self.theme.navy, bold=active, align=PP_ALIGN.CENTER)
            # The labels define positive current directions, even before any
            # numerical current has been revealed in the worked example.
            direction = "←" if n in (4, 5) else "→"
            self._text(slide, f"I{n} {direction}", center-.48, yy+.13, .96, .32,
                       font, self.theme.accent if active else self.theme.navy,
                       bold=active, align=PP_ALIGN.CENTER)

        def source(x1, x2, yy, n, positive_right):
            center = (x1+x2)/2
            span = min(.17, (x2-x1)*.13)
            line(x1, yy, center-span, yy)
            line(center+span, yy, x2, yy)
            for pos, length in ((center-span, .16 if positive_right else .32), (center+span, .32 if positive_right else .16)):
                line(pos, yy-length/2, pos, yy+length/2)
            label_y = yy-.60 if yy == top else yy+.38
            self._text(slide, f"E{n}  {element.sources[n-1]:g} V", center-.55, label_y, 1.1, .34, font, self.theme.text, align=PP_ALIGN.CENTER)
            polarity_y = yy+.16 if yy == top else yy+.06
            self._text(slide, "−" if positive_right else "+", center-span-.19, polarity_y,
                       .38, .30, font, self.theme.navy, bold=True, align=PP_ALIGN.CENTER)
            self._text(slide, "+" if positive_right else "−", center+span-.19, polarity_y,
                       .38, .30, font, self.theme.navy, bold=True, align=PP_ALIGN.CENTER)

        line(left, mid, left, bottom)
        line(a, top, a, mid)
        line(c, mid, c, bottom)
        line(d, top, d, mid)
        line(b, mid, b, bottom)
        resistor(left, a, mid, 1)
        resistor(a, c, mid, 2)
        resistor(c, d, mid, 4)
        resistor(d, b, mid, 5)
        resistor(a, a+w*.23, top, 3)
        source(a+w*.23, d, top, 3, True)
        source(left, c, bottom, 1, False)
        source(c, b, bottom, 2, True)
        for xx, yy, label in ((left, mid, "L"), (a, mid, "A"), (c, mid, "C (0 V)"), (d, mid, "D"), (b, mid, "B")):
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(xx-.035), Inches(yy-.035), Inches(.07), Inches(.07))
            dot.fill.solid(); dot.fill.fore_color.rgb = _rgb(self.theme.navy)
            dot.line.fill.background()
            if label == "L":
                label_x, label_w = xx-.52, .34
            elif label == "B":
                label_x, label_w = xx+.10, .34
            elif label == "C (0 V)":
                label_x, label_w = xx+.12, .90
            else:
                label_x, label_w = xx-.42, .84
            self._text(slide, label, label_x, yy+.47, label_w, .24, 13,
                       self.theme.navy, align=PP_ALIGN.CENTER)
        group = slide.shapes.add_group_shape(list(slide.shapes)[before:])
        group.name = element.id
        self._set_alt_text(group, "Tre ideelle DC-kilder: E1 plus mod venstre, E2 og E3 plus mod højre. Positive strømretninger: I1 L til A, I2 A til C, I3 A til D, I4 D til C, I5 B til D. L er knuden før R1, B er knuden efter R5. Reference C = 0 V.")

    def _process(self, slide, element):
        before = len(slide.shapes)
        count = len(element.steps); available = 11.8; gap = .18
        width = (available - gap * (count - 1)) / count
        for index, step in enumerate(element.steps):
            x = .75 + index * (width + gap)
            self._rect(slide, x, 2.25, width, 2.0, self.theme.pale, radius=True, name=element.id if index == 0 else None)
            self._text(slide, f"{index+1}\n{step}", x+.12, 2.52, width-.24, 1.45, 20, self.theme.navy, bold=True, align=PP_ALIGN.CENTER)

        slide.shapes.add_group_shape(list(slide.shapes)[before:]).name = element.id

    def _table(self, slide, element, x, y, w, h):
        rows, cols = len(element.rows)+1, len(element.headers)
        shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = element.id
        table = shape.table
        for c, header in enumerate(element.headers):
            table.cell(0, c).text = header
        for r, row in enumerate(element.rows, 1):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        for r in range(rows):
            for c in range(cols):
                cell = table.cell(r, c)
                cell.fill.solid(); cell.fill.fore_color.rgb = _rgb(self.theme.navy if r == 0 else "FFFFFF")
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(15 if rows > 7 else 18)
                    p.font.color.rgb = _rgb("FFFFFF" if r == 0 else self.theme.text)
                    p.font.bold = r == 0

    def _text(self, slide, value, x, y, w, h, size, color, bold=False, align=PP_ALIGN.LEFT, name=None):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name: shape.name = name
        frame = shape.text_frame; frame.clear(); frame.word_wrap = True; frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = frame.paragraphs[0]; p.text = value; p.alignment = align
        p.font.name = self.theme.font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = _rgb(color)
        return shape

    def _rect(self, slide, x, y, w, h, color, radius=False, name=None):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(color); shape.line.fill.background()
        if name: shape.name = name
        return shape

    @staticmethod
    def _set_alt_text(shape, text):
        props = shape._element.xpath(".//*[local-name()='cNvPr']")
        if props: props[0].set("descr", text)
