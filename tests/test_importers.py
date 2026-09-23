from pathlib import Path

from docx import Document as DocxDocument
from matplotlib import pyplot as plt
from pypdf import PdfWriter

from powerpoint_app.importers import import_docx, import_markdown, import_pdf


def test_markdown_preserves_structure_and_locator(tmp_path: Path):
    source = tmp_path / "a.md"; source.write_text("# Pumpe\n\nTryk h′ og h″.\n", encoding="utf-8")
    doc = import_markdown(source, "src1")
    assert [b.kind for b in doc.blocks] == ["heading", "paragraph"]
    assert "h′" in doc.blocks[1].text
    assert doc.blocks[1].locator


def test_docx_imports_paragraphs_and_tables(tmp_path: Path):
    source = tmp_path / "a.docx"; word = DocxDocument(); word.add_heading("Pumper", 1); word.add_paragraph("Tryk h′ og h″")
    table = word.add_table(rows=2, cols=2); table.cell(0, 0).text = "Navn"; table.cell(1, 0).text = "P1"; word.save(source)
    doc = import_docx(source, "src-docx")
    assert doc.title == "Pumper" and any(block.kind == "table" for block in doc.blocks)
    assert any("h′" in block.text for block in doc.blocks)


def test_text_pdf_and_scanned_pdf_warning(tmp_path: Path):
    text_pdf = tmp_path / "text.pdf"; fig = plt.figure(); fig.text(.1, .5, "Pumpe og tryk 10 bar"); fig.savefig(text_pdf); plt.close(fig)
    doc = import_pdf(text_pdf, "src-pdf")
    assert any("Pumpe" in block.text for block in doc.blocks)
    scan_pdf = tmp_path / "scan.pdf"; writer = PdfWriter(); writer.add_blank_page(width=600, height=800)
    with scan_pdf.open("wb") as handle: writer.write(handle)
    scanned = import_pdf(scan_pdf, "src-scan")
    assert any("OCR" in warning for warning in scanned.warnings)
