"""Compact A4 lookup pages using the existing typed slide elements and assets."""
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from powerpoint_app.domain.models import FormulaElement, TableElement, TextElement, CircuitElement

SEARCH, GIVEN, CALCULATE = '075B70', '855515', '286044'

def text(r, slide, value, x, y, w, h, size=11, color=None, bold=False, name=None):
    sh = r._text(slide, value, x, y, w+.05, h, size, color or r.theme.text, bold=bold, name=name)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for p in tf.paragraphs:
        p.space_before = p.space_after = Pt(0)
    from powerpoint_app.rendering.math_markup import serif_current_runs
    serif_current_runs(tf)
    return sh

def reference_page(r, spec, number, total, plan):
    slide = r.prs.slides.add_slide(r.prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
    text(r, slide, spec.title, .45, .3, 10.75, .42, 22, SEARCH, True)
    text(r, slide, 'SØGER', .45, .78, .7, .2, 9, SEARCH, True)
    text(r, slide, 'HAR', 1.2, .78, .55, .2, 9, GIVEN, True)
    text(r, slide, 'BEREGN', 1.8, .78, 1, .2, 9, CALCULATE, True)
    text(r, slide, f'{number} / {total}', 10.3, 7.98, .9, .2, 9)
    text(r, slide, 'EL · BM4 · kildebaseret opslagsværk', .45, 7.98, 6, .2, 9)
    if spec.layout == 'lookup_cards':
        rich=any(isinstance(e,TextElement) and e.math_markup for e in spec.elements)
        if rich and len(spec.elements)>6:raise ValueError('Læsbar matematik kræver højst to kort pr. side.')
        for i in range(0, len(spec.elements), 3):
            header, formula, detail = spec.elements[i:i+3]
            x=.45+(i//3%2)*5.6; y=1.08+(i//6)*3.42; w=5.2
            rows=header.text.split('\n')
            text(r, slide, rows[0], x, y, w, .46, 14, SEARCH, True, header.id)
            text(r, slide, rows[1], x, y+.49, w, .47, 11, GIVEN)
            text(r, slide, rows[2], x, y+.99, w, .44, 10.5)
            if rich:
                from powerpoint_app.rendering.rich_math import math_picture,rich_details
                height=math_picture(r,slide,formula.latex,x,y+1.48,w,26,formula.id)
                detail_y=y+1.48+height+.16
                rich_details(r,slide,detail,x,detail_y,w,7.72-detail_y)
                r._rect(slide,x,7.78,w,.015,'CCD8DC')
                continue
            tall = formula.latex.count(r'\sum') >= 2 and len(detail.text.splitlines()) <= 3
            r._element(slide, formula, x, y+(1.38 if tall else 1.44), w, .85 if tall else .66)
            text(r, slide, detail.text, x, y+(2.28 if tall else 2.16), w, .94 if tall else 1.06, 10.5, name=detail.id)
            r._rect(slide, x, y+3.29, w, .015, 'CCD8DC')
    elif spec.layout == 'lookup_table':
        table=next(e for e in spec.elements if isinstance(e, TableElement))
        reference_table(r, slide, table)
        for e in spec.elements:
            if isinstance(e, TextElement): text(r, slide, e.text, .45, 7.35, 10.7, .5, 10, name=e.id)
    elif spec.layout == 'lookup_nodal':
        nodal_page(r, slide, spec)
    else:
        for i, element in enumerate(spec.elements):
            if isinstance(element, TextElement):
                text(r, slide, element.text, .45, 1.12+i*.57, 10.7, .52, 12, name=element.id)
    from powerpoint_app.rendering.navigation import navigation_shapes
    navigation_shapes(r, slide, spec)
    slide.notes_slide.notes_text_frame.text=spec.speaker_notes

def reference_table(r, slide, e):
    rows, cols = len(e.rows)+1, len(e.headers)
    x,y,w,h=.45,1.14,10.7,6.03
    is_index = e.headers == ['Søger','Har opgivet','Situation / betingelse','Formel','Side']
    if is_index:
        h = .4 + .563 * len(e.rows)
    shape=slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h));shape.name=e.id
    table=shape.table
    weights=([1.65,2.25,2.55,3.7,.55] if e.headers==['Søger','Har opgivet','Situation / betingelse','Formel','Side'] else [max(8,min(35,max(len(str(row[c])) for row in [e.headers]+e.rows))) for c in range(cols)])
    widths=[w*v/sum(weights) for v in weights]
    for c, width in enumerate(widths):table.columns[c].width=Inches(width)
    table.rows[0].height=Inches(.4)
    rh=(h-.4)/(rows-1)
    for j in range(1,rows):table.rows[j].height=Inches(rh)
    for j,row in enumerate([e.headers]+e.rows):
        for c,value in enumerate(row):
            cell=table.cell(j,c);cell.margin_left=cell.margin_right=Inches(.045);cell.margin_top=cell.margin_bottom=Inches(.05)
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(SEARCH if j==0 else ('F0F5F6' if j%2 else 'FFFFFF'))
            is_formula=j>0 and e.headers[c]=='Formel'
            cell.text='' if is_formula else str(value)
            for p in cell.text_frame.paragraphs:
                p.font.name=r.theme.font;p.font.size=Pt(10.5 if is_index or rows>8 else 12);p.font.bold=j==0
                p.font.color.rgb=RGBColor.from_string('FFFFFF' if j==0 else SEARCH if c==0 else GIVEN if c==1 and e.headers[0]=='Søger' else r.theme.text)
                p.space_before=p.space_after=Pt(0)
            from powerpoint_app.rendering.math_markup import serif_current_runs
            serif_current_runs(cell.text_frame)
            if is_formula:
                formula=FormulaElement(id=f'{e.id}-f{j}',type='formula',latex=value)
                r._element(slide,formula,x+sum(widths[:c])+.05,y+.4+(j-1)*rh+.04,widths[c]-.1,rh-.08)

def nodal_page(r, slide, spec):
    from powerpoint_app.rendering.rich_math import math_picture,rich_details,wrapped_text
    e=next(a for a in spec.elements if isinstance(a,CircuitElement))
    sources=e.branch_sources
    if sources is None: raise ValueError('Knudepunktsfiguren kræver grenkilder.')
    before=len(slide.shapes);xs=[.9+i*1.7 for i in range(len(sources))];top,bottom=1.72,5.85
    def line(x1,y1,x2,y2,arrow=False):
        a=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));a.line.color.rgb=RGBColor.from_string(SEARCH);a.line.width=Pt(1.5)
        if arrow:
            head=OxmlElement('a:tailEnd');head.set('type','triangle');a._element.spPr.get_or_add_ln().append(head)
        return a
    line(xs[0],top,xs[-1],top);line(xs[0],bottom,xs[-1],bottom)
    text(r,slide,'a: Va',.6,1.15,1.7,.4,15,SEARCH,True)
    text(r,slide,'0: reference = 0 V',.6,6.02,4.4,.4,14,SEARCH,True)
    for i,(x,resistance,emf) in enumerate(zip(xs,e.resistances,sources),1):
        line(x,top,x,2.7);line(x+.35,1.95,x+.35,2.52,True)
        math_picture(r,slide,f'I_{{{i}}}',x+.45,2.1,.6,18,f'nodal-label-I{i}')
        a=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x-.13),Inches(2.7),Inches(.26),Inches(.75));a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string('FFFFFF');a.line.color.rgb=RGBColor.from_string(SEARCH)
        text(r,slide,f'R{i}\n{resistance:g} Ω',x+.25,2.75,1.25,.7,13)
        if emf:
            line(x,3.45,x,4.22);line(x,4.92,x,bottom)
            a=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-.35),Inches(4.22),Inches(.7),Inches(.7));a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string('FFFFFF');a.line.color.rgb=RGBColor.from_string(SEARCH)
            text(r,slide,'+\n−',x-.11,4.23,.23,.67,15,SEARCH,True)
            text(r,slide,f'E{i}\n{emf:g} V',x+.4,4.28,1.1,.7,13)
        else:line(x,3.45,x,bottom)
    slide.shapes.add_group_shape(list(slide.shapes)[before:]).name=e.id
    y=1.18
    for a in spec.elements:
        if isinstance(a,FormulaElement):
            y+=math_picture(r,slide,a.latex,6.0,y,5.18,20,a.id)+.16
        elif isinstance(a,TextElement):
            if a.math_markup:y+=rich_details(r,slide,a,6.0,y,5.18,6.5-y)+.12
            else:
                value,height=wrapped_text(a.text,5.18)
                text(r,slide,value,6.0,y,5.18,height,12,name=a.id);y+=height+.12
    text(r,slide,'Alle pile går ud af a. Minus betyder strøm mod pilen. + og − angiver kildernes faste polaritet.',.6,6.7,10.5,.7,12)
