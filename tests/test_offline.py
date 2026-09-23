import json
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from powerpoint_app.cli import run
from powerpoint_app.planning.providers import JsonApiProvider, OllamaProvider


def test_export_and_theme_change_never_call_provider(tmp_path: Path, monkeypatch):
    root = tmp_path / "project"; shutil.copytree(Path("examples/demo_project"), root)
    def forbidden(*args, **kwargs): raise AssertionError("provider blev kaldt")
    monkeypatch.setattr(OllamaProvider, "complete", forbidden); monkeypatch.setattr(JsonApiProvider, "complete", forbidden)
    assert run(["export", str(root), "--output", "default.pptx"]) == 0
    custom = {"id": "test", "background": "FFF4E6", "navy": "203040", "accent": "CC5500", "pale": "FFE0B2", "warning": "FFAA00", "text": "101820", "font": "Aptos"}
    theme = tmp_path / "theme.json"; theme.write_text(json.dumps(custom), encoding="utf-8")
    assert run(["export", str(root), "--output", "changed.pptx", "--theme", str(theme)]) == 0
    first = Presentation(root / "exports" / "default.pptx"); second = Presentation(root / "exports" / "changed.pptx")
    assert first.slides[0].background.fill.fore_color.rgb != second.slides[0].background.fill.fore_color.rgb


def test_com_failure_keeps_static_export(tmp_path: Path, monkeypatch):
    root = tmp_path / "project"; shutil.copytree(Path("examples/demo_project"), root)
    monkeypatch.setattr("powerpoint_app.office.com.platform.system", lambda: "Linux")
    with pytest.raises(RuntimeError, match="Windows"):
        run(["export", str(root), "--output", "static.pptx", "--animate"])
    assert (root / "exports" / "static.pptx").is_file()
