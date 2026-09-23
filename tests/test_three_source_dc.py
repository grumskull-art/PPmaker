"""Regression checks against the original BM4 exam topology and values."""
from pathlib import Path

import pytest
from pptx import Presentation

from powerpoint_app.domain.models import SlidePlan, ThreeSourceCircuitElement
from powerpoint_app.projects import load_plan
from powerpoint_app.quality import inspect_plan, inspect_pptx_geometry
from powerpoint_app.quality.three_source_dc import solve_three_source_dc
from powerpoint_app.rendering import PptxRenderer
from powerpoint_app.rendering.teaching import presentation_frames


ROOT = Path('examples/bm4_exam_2023')


def test_exam_answers_and_source_power_balance():
    diagram = next(e for e in load_plan(ROOT / 'slide-plan.json').slides[0].elements
                   if isinstance(e, ThreeSourceCircuitElement))
    result = solve_three_source_dc(diagram)
    assert result.currents == pytest.approx((1.5107084, .4530478, 1.0576606, 2.1367381, 1.0790774), rel=1e-6)
    assert result.vab == pytest.approx(-17.76771005)
    assert result.total_power == pytest.approx(134.34925865)
    assert result.source_power == pytest.approx(result.total_power)
    altered = diagram.model_copy(update={'resistances': (25., 27., 15., 10., 81.)})
    assert solve_three_source_dc(altered).total_power != pytest.approx(result.total_power)


def test_exam_plan_binds_calculations_to_edited_diagram():
    plan = load_plan(ROOT / 'slide-plan.json')
    assert not [f for f in inspect_plan(plan, ROOT) if f.level == 'error']
    data = plan.model_dump()
    next(e for e in data['slides'][5]['elements'] if e['type'] == 'three_source_dc')['resistances'] = (25, 27, 15, 10, 81)
    edited = SlidePlan.model_validate(data)
    findings = inspect_plan(edited, ROOT)
    assert any(f.level == 'error' and 'R5' in f.message for f in findings)
    data = plan.model_dump()
    data['slides'][5]['elements'][5]['branch_index'] = 6
    with pytest.raises(ValueError):
        SlidePlan.model_validate(data)


def test_exam_export_reveals_five_branches_and_delays_answer(tmp_path):
    plan = load_plan(ROOT / 'slide-plan.json')
    frames = presentation_frames(plan)
    path = PptxRenderer(ROOT).render(plan, tmp_path / 'exam.pptx')
    assert len(frames) == 29
    assert inspect_pptx_geometry(path) == []
    prs = Presentation(path)
    for i, frame in enumerate(frames):
        if frame.slide.id == 's7':
            visible = '\n'.join(getattr(s, 'text', '') for s in prs.slides[i].shapes)
            assert ('U_AB = -17,77 V.' in visible) == frame.show_answer
    assert len([f for f in frames if f.slide.id == 's6']) == 6
    first = next(s for s in prs.slides[0].shapes if s.name == 'c1')
    labels = [getattr(s, 'text', '') for s in first.shapes]
    assert {'I1 →', 'I2 →', 'I3 →', 'I4 ←', 'I5 ←', 'L', 'B'}.issubset(labels)
    assert labels.count('+') == 3 and labels.count('−') == 3
    explanation = '\n'.join(getattr(s, 'text', '') for s in prs.slides[1].shapes)
    assert 'L er knuden før R1' in explanation and 'B er knuden efter R5' in explanation
