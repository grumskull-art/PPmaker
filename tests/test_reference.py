from dataclasses import replace
from pathlib import Path
import pytest
from pptx import Presentation
from pydantic import ValidationError
from powerpoint_app.domain.models import SlidePlan, FormulaElement
from powerpoint_app.projects import load_plan
from powerpoint_app.rendering import PptxRenderer
from powerpoint_app.rendering.theme import DEFAULT_THEME
from powerpoint_app.visuals.assets import formula_png
from powerpoint_app.quality.checks import inspect_pptx_geometry


def lookup_data():
    data=load_plan(Path('examples/demo_project/slide-plan.json')).model_dump()
    data['schema_version']='1.1';data['deck']['page_format']='a4_landscape'
    data['slides']=[dict(id='lookup',layout='lookup_cards',title='Strøm',elements=[
        dict(id='h',type='text',text='Søger: I [A]\nHar: U [V], R [Ω]\nGælder: ohmsk DC'),
        dict(id='f',type='formula',latex=r'I=\frac{U}{R}'),
        dict(id='d',type='text',text='24 V / 12 Ω = 2 A')],speaker_notes='Kildeeksempel')]
    return data


def test_lookup_uses_a4_and_preserves_native_text_and_formula_source(tmp_path):
    output=tmp_path/'lookup.pptx';plan=SlidePlan.model_validate(lookup_data())
    PptxRenderer(tmp_path).render(plan,output)
    deck=Presentation(output)
    assert abs(deck.slide_width/914400-297/25.4)<.001
    assert abs(deck.slide_height/914400-210/25.4)<.001
    assert inspect_pptx_geometry(output)==[]
    assert any('24 V / 12 Ω' in getattr(s,'text','') for s in deck.slides[0].shapes)
    f=next(s for s in deck.slides[0].shapes if s.name=='f')
    assert f._element.xpath('.//p:cNvPr')[0].get('descr')==r'I=\frac{U}{R}'


def test_lookup_rejects_incomplete_card_and_wrong_page_format():
    data=lookup_data();data['slides'][0]['elements'].pop()
    with pytest.raises(ValidationError):SlidePlan.model_validate(data)
    data=lookup_data();data['deck']['page_format']='wide'
    with pytest.raises(ValidationError):SlidePlan.model_validate(data)


def test_lookup_rejects_missing_header_lines_and_unrenderable_tables():
    data=lookup_data();data['slides'][0]['elements'][0]['text']='Kun én linje'
    with pytest.raises(ValidationError):SlidePlan.model_validate(data)
    data=lookup_data();data['slides'][0]['layout']='lookup_table'
    data['slides'][0]['elements']=[dict(id='t',type='table',headers=['A'],rows=[['x']]*11)]
    with pytest.raises(ValidationError):SlidePlan.model_validate(data)


def test_lookup_nodal_rejects_four_branches():
    data=lookup_data();data['slides'][0]['layout']='lookup_nodal'
    data['slides'][0]['elements']=[dict(id='c',type='parallel_circuit',voltage=24,resistances=[1,2,3,4],branch_sources=[24,0,0,0])]
    with pytest.raises(ValidationError):SlidePlan.model_validate(data)


def test_formula_palette_changes_cache_asset(tmp_path):
    f=FormulaElement(id='f',type='formula',latex=r'I=\frac{U}{R}')
    blue=formula_png(f,tmp_path,DEFAULT_THEME)
    green=formula_png(f,tmp_path,replace(DEFAULT_THEME,navy='286044'))
    assert blue!=green
    assert blue.read_bytes()!=green.read_bytes()
