import shutil
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from powerpoint_app.projects import load_plan
from powerpoint_app.quality import inspect_plan, inspect_pptx_geometry
from powerpoint_app.rendering import PptxRenderer


def test_demo_renders_with_notes_formula_and_figure(tmp_path: Path):
    source = Path("examples/demo_project")
    root = tmp_path / "demo"; shutil.copytree(source, root)
    (root / "cache").mkdir(exist_ok=True); (root / "exports").mkdir(exist_ok=True)
    plan = load_plan(root / "slide-plan.json")
    output = root / "exports" / "demo.pptx"
    PptxRenderer(root).render(plan, output)
    prs = Presentation(output)
    assert len(prs.slides) == 7
    assert "Sæt scenen" in prs.slides[0].notes_slide.notes_text_frame.text
    assert any(shape.name == "e43" for shape in prs.slides[3].shapes)
    assert any(shape.name == "e31" for shape in prs.slides[2].shapes)
    assert output.stat().st_size > 20_000
    assert inspect_pptx_geometry(output) == []
    assert PptxRenderer(root).template_path == Path("templates/martec_inspired.pptx").resolve()


def test_quality_flags_missing_asset(tmp_path: Path):
    plan = load_plan(Path("examples/demo_project/slide-plan.json"))
    data = plan.model_dump(); data["slides"][0]["elements"] = [{"id": "img", "type": "image", "asset": "assets/no.png", "alt_text": "mangler", "source_ids": []}]
    plan = type(plan).model_validate(data)
    assert any(f.level == "error" for f in inspect_plan(plan, tmp_path))


def test_geometry_finds_shape_outside_slide(tmp_path: Path):
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.shapes.add_textbox(prs.slide_width - Inches(1), 0, Inches(2), Inches(1))
    path = tmp_path / "outside.pptx"; prs.save(path)
    assert any(f.level == "error" for f in inspect_pptx_geometry(path))


def test_same_plan_and_theme_have_same_content_and_geometry(tmp_path: Path):
    root = Path("examples/demo_project"); plan = load_plan(root / "slide-plan.json")
    first, second = tmp_path / "a.pptx", tmp_path / "b.pptx"; PptxRenderer(root).render(plan, first); PptxRenderer(root).render(plan, second)
    def signature(path):
        prs = Presentation(path)
        return [[(shape.name, shape.left, shape.top, shape.width, shape.height, getattr(shape, "text", "")) for shape in slide.shapes] for slide in prs.slides]
    assert signature(first) == signature(second)
