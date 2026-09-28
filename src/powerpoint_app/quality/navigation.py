"""Check native PPTX relationships and exported PDF destinations against the plan."""
from collections import Counter
from pptx import Presentation
from pypdf import PdfReader


def audit_navigation(plan, catalog, pptx_path, pdf_path=None):
    deck=Presentation(pptx_path)
    if len(deck.slides)!=len(plan.slides):raise ValueError('Navigationsaudit kræver statisk eksport af samme plan.')
    positions={s.id:i for i,s in enumerate(plan.slides)}
    slide_ids={s.slide_id:i for i,s in enumerate(deck.slides)}
    entries={e['id']:e for e in catalog['entries']};row_counts=Counter();total=0
    pdf=PdfReader(pdf_path) if pdf_path else None
    page_refs={(p.indirect_reference.idnum,p.indirect_reference.generation):i for i,p in enumerate(pdf.pages)} if pdf else {}
    if pdf and len(pdf.pages)!=len(plan.slides):raise ValueError('PDF og plan har forskellige sidetal.')
    for index,(spec,slide) in enumerate(zip(plan.slides,deck.slides)):
        actual={s.name:s for s in slide.shapes if s.name.startswith('nav-')}
        if set(actual)!={'nav-'+n.id for n in spec.navigation}:raise ValueError('Manglende eller ekstra links: '+spec.id)
        pdf_links=[a.get_object() for a in pdf.pages[index].get('/Annots',[]) if a.get_object().get('/Subtype')=='/Link'] if pdf else []
        if pdf and len(pdf_links)!=len(spec.navigation):raise ValueError('PDF-linkantal afviger: '+spec.id)
        for link in spec.navigation:
            total+=1;shape=actual['nav-'+link.id];expected=positions[link.target_slide_id]
            if shape.click_action.target_slide is None or slide_ids[shape.click_action.target_slide.slide_id]!=expected:
                raise ValueError('Forkert PPTX-mål: '+link.id)
            if link.placement=='row':
                e=entries[link.entry_id];row_counts[e['id']]+=1
                table=next(el for el in spec.elements if el.id==link.table_id)
                if not table.rows[link.row_index][0].startswith(e['id']+' · '):raise ValueError('Kode/række afviger: '+e['id'])
                if e['page']!=expected+1 or int(table.rows[link.row_index][-1])!=expected+1:raise ValueError('Sidetal afviger: '+e['id'])
                if not any(el.id==e['id']+'-formula' for el in plan.slides[expected].elements):raise ValueError('Målet har ikke formelkortet: '+e['id'])
                if not any(sh.name==e['id']+'-formula' for sh in deck.slides[expected].shapes):raise ValueError('Formelkort mangler i PPTX: '+e['id'])
                if pdf and e['id'] not in pdf.pages[expected].extract_text():raise ValueError('Formelkode ikke søgbar i PDF: '+e['id'])
            if pdf:
                x,y,w,h=[v/12700 for v in (shape.left,shape.top,shape.width,shape.height)]
                page_h=float(pdf.pages[index].mediabox.height);rectangle=[x,page_h-y-h,x+w,page_h-y]
                center=((rectangle[0]+rectangle[2])/2,(rectangle[1]+rectangle[3])/2)
                candidates=[a for a in pdf_links if float(a['/Rect'][0])<=center[0]<=float(a['/Rect'][2]) and float(a['/Rect'][1])<=center[1]<=float(a['/Rect'][3])]
                if link.placement=='row':
                    candidates=[a for a in candidates if all(abs(float(v)-s)<2 for v,s in zip(a['/Rect'],rectangle))]
                if len(candidates)!=1:raise ValueError('PDF-linkområdet afviger: '+link.id)
                a=candidates[0];dest=a.get('/Dest',a.get('/A',{}).get('/D'))
                if not isinstance(dest,list):raise ValueError('PDF har ikke internt mål: '+link.id)
                target=dest[0];actual_page=page_refs.get((target.idnum,target.generation))
                if actual_page!=expected:raise ValueError('Forkert PDF-mål: '+link.id)
    if row_counts!=Counter({id:1 for id in entries}):raise ValueError('Alle databaseopslag skal have præcis én indeksrække.')
    return {'pages':len(deck.slides),'index_entries':len(entries),'pptx_internal_links':total,'pdf_internal_links':total if pdf else None,'each_entry_has_one_correct_target':True}
