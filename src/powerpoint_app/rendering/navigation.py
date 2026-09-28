"""Native internal slide links resolved after every destination exists."""
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from powerpoint_app.rendering.reference import text, SEARCH


def navigation_shapes(renderer, slide, spec):
    menu = footer = 0
    for link in spec.navigation:
        if link.placement == 'menu':
            x=.45+(menu%2)*5.6; y=2.27+(menu//2)*1.05
            shape=text(renderer, slide, link.label, x+.18, y+.12, 4.75, .66, 17, SEARCH, True)
            shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('F0F5F6')
            # A native rectangle makes the entire shortcut clickable, including empty space.
            shape._element.nvSpPr.cNvSpPr.attrib.pop('txBox', None)
            shape.line.fill.background()
            label,_,description=link.label.partition('\n')
            p=shape.text_frame.paragraphs[0];p.clear()
            p.add_run().text=label;p._p.append(OxmlElement('a:br'))
            run=p.add_run();run.text=description;run.font.size=Pt(11);run.font.bold=False
            menu+=1
        elif link.placement == 'footer':
            shape=text(renderer, slide, link.label, 6.5+footer*1.85, 7.98, 1.75, .2, 9, SEARCH)
            footer+=1
        else:
            table=next(s for s in slide.shapes if s.name==link.table_id)
            row=link.row_index+1
            top=table.top+sum(table.table.rows[i].height for i in range(row))
            shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,table.left,top,table.width,table.table.rows[row].height)
            shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('FFFFFF')
            alpha=OxmlElement('a:alpha');alpha.set('val','0')
            shape._element.spPr.solidFill[0].append(alpha)
            shape.line.fill.background()
        shape.name='nav-'+link.id
        for effect in shape._element.xpath('.//a:effectRef'):effect.set('idx','0')
        shape._element.xpath('.//p:cNvPr')[0].set('descr',link.label)


def resolve_links(rendered):
    destinations={}
    for spec,slide in rendered:
        destinations.setdefault(spec.id,slide)
    for spec,slide in rendered:
        for link in spec.navigation:
            shape=next(s for s in slide.shapes if s.name=='nav-'+link.id)
            shape.click_action.target_slide=destinations[link.target_slide_id]
            shape._element.xpath('.//a:hlinkClick')[0].set('tooltip',link.label)
