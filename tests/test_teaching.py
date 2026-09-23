import json
from pathlib import Path

import pytest
from pptx import Presentation

from powerpoint_app.cli import run
from powerpoint_app.domain.models import SlidePlan
from powerpoint_app.domain.teaching import CalculationCheck, Question, TeachingProfile
from powerpoint_app.planning.teaching import load_profile, save_profile, teaching_brief
from powerpoint_app.projects import create_project, load_plan
from powerpoint_app.quality import inspect_plan, inspect_pptx_geometry
from powerpoint_app.quality.calculations import verify_calculation
from powerpoint_app.rendering import PptxRenderer
from powerpoint_app.rendering.teaching import presentation_frames

ROOT = Path('examples/teaching_dc')


def pilot():
    return load_plan(ROOT / 'slide-plan.json')


def check(expression='U/R', expected=2, unit='A', resistance_unit='ohm'):
    return CalculationCheck(element_id='f', expression=expression,
        quantities={'U': {'value': 24, 'unit': 'V'}, 'R': {'value': 12, 'unit': resistance_unit}},
        expected={'value': expected, 'unit': unit})


@pytest.mark.parametrize('spec,status', [
    (check(), 'passed'), (check(expected=3), 'failed'), (check(unit='V'), 'failed'),
    (check(expression='U+R'), 'failed'), (check(expression='U/0'), 'failed'),
    (check(resistance_unit='unknown'), 'unsupported'),
    (check(expression="__import__('os').system('echo bad')"), 'unsupported'),
    (check(expression='U**100000000'), 'unsupported'),
    (check(expected=.002, resistance_unit='kohm'), 'passed'),
])
def test_independent_numeric_and_dimension_checks(spec, status):
    assert verify_calculation(spec).status == status


def test_teaching_schema_roundtrip_and_reference_validation():
    plan = pilot()
    assert SlidePlan.model_validate_json(plan.model_dump_json()) == plan
    data = plan.model_dump(); data['slides'][0]['teaching']['objective_indices'] = [99]
    with pytest.raises(ValueError, match='læringsmål'): SlidePlan.model_validate(data)
    data = plan.model_dump(); data['slides'][0]['animations'] = [{'target_id':'f31', 'order':1}]
    with pytest.raises(ValueError, match='ukendt element'): SlidePlan.model_validate(data)
    data = plan.model_dump(); data['schema_version'] = '1.0'
    with pytest.raises(ValueError, match='1.1'): SlidePlan.model_validate(data)


def test_concept_question_rejects_missing_or_invalid_answer_index():
    base = dict(prompt='Hvorfor?', answer='Fordi.', explanation='Forklaring.', options=['A', 'B', 'C'])
    with pytest.raises(ValueError, match='svarindeks'): Question(**base)
    with pytest.raises(ValueError, match='svarindeks'): Question(**base, correct_option=3)
    with pytest.raises(ValueError, match='forskellige'): Question(**{**base, 'options': ['A', 'A', 'C']}, correct_option=0)


def test_bad_calculation_blocks_export(tmp_path):
    plan = pilot(); plan.slides[2].calculation_checks[0].expected.value = 3
    errors = inspect_plan(plan, ROOT)
    assert any(f.level == 'error' and 'f31' in f.message for f in errors)
    create_project(tmp_path)
    (tmp_path/'slide-plan.json').write_text(plan.model_dump_json())
    with pytest.raises(ValueError, match='Eksport stoppet'): run(['export', str(tmp_path)])


def test_profiles_persist_and_brief_overrides(tmp_path):
    profile = pilot().teaching_profile
    save_profile(tmp_path, profile)
    assert load_profile(tmp_path) == profile
    brief = teaching_brief(tmp_path, {'topic':'DC', 'duration_minutes':8})
    assert brief['duration_minutes'] == 8
    assert brief['teaching_profile']['notation'] == profile.notation
    assert brief['approx_slides'] == 7


def test_frames_are_pure_and_questions_hide_answers(tmp_path):
    plan = pilot(); original = plan.model_dump_json()
    frames = presentation_frames(plan)
    assert len(frames) == 14 and len(presentation_frames(plan, 'study')) == 7
    path = PptxRenderer(ROOT).render(plan, tmp_path/'lesson.pptx')
    assert plan.model_dump_json() == original
    assert inspect_pptx_geometry(path) == []
    prs = Presentation(path)
    question_indexes = [i for i, f in enumerate(frames) if f.slide.teaching.question]
    for i in question_indexes:
        frame = frames[i]
        visible = '\n'.join(getattr(s, 'text', '') for s in prs.slides[i].shapes)
        assert (frame.slide.teaching.question.answer in visible) == frame.show_answer
    concept = next(i for i, f in enumerate(frames) if f.slide.id == 's6' and not f.show_answer)
    visible = '\n'.join(getattr(s, 'text', '') for s in prs.slides[concept].shapes)
    assert 'A. Samme knuder' in visible and 'B. Samme modstande' in visible
    assert 'De deler endeknuder.' not in visible
    notes = prs.slides[concept].notes_slide.notes_text_frame.text
    assert 'Samtale efter individuelt svar' in notes and 'Korrekt valg: A.' in notes
    worked = [i for i, f in enumerate(frames) if f.slide.id == 's3']
    assert len(worked) == 4
    for index, count in zip(worked, range(4)):
        names = {s.name for s in prs.slides[index].shapes}
        assert sum(name in names for name in ('f31','f32','f33')) == count
    positions = [(s.left,s.top,s.width,s.height) for i in worked for s in prs.slides[i].shapes if s.name == 'c3']
    assert len(set(positions)) == 1


def test_prompt_is_offline_and_uses_saved_profile(tmp_path, monkeypatch):
    import shutil
    from powerpoint_app.planning.providers import JsonApiProvider, OllamaProvider
    def forbidden(*args, **kwargs): raise AssertionError('Network call')
    monkeypatch.setattr(JsonApiProvider, 'complete', forbidden)
    monkeypatch.setattr(OllamaProvider, 'complete', forbidden)
    root = tmp_path/'pilot'; shutil.copytree(ROOT,root)
    assert run(['prompt', str(root), '--topic','DC', '--audience','BM4']) == 0
    prompt = (root/'planning-prompt.txt').read_text()
    assert 'Knudepotentialer' in prompt and '2026-09-23.3' in prompt
    assert 'BM4_EL-TEK1' in prompt and 'ikke citerede studieordningskrav' in prompt
    assert '"duration_minutes": 12' in prompt
    assert run(['export',str(root),'--output','offline.pptx']) == 0


def test_failed_checks_never_claim_passed_for_unsupported_units():
    result = verify_calculation(check(resistance_unit='degC'))
    assert result.status == 'unsupported'


def test_formula_image_keeps_aspect_ratio(tmp_path):
    from PIL import Image
    path = PptxRenderer(ROOT).render(pilot(), tmp_path/'study.pptx', mode='study')
    formula = next(s for s in Presentation(path).slides[2].shapes if s.name == 'f31')
    import io
    with Image.open(io.BytesIO(formula.image.blob)) as image:
        assert formula.width / formula.height == pytest.approx(image.width/image.height, rel=1e-5)


def test_provider_cannot_silently_drop_profile_and_cache_tracks_changes(tmp_path):
    from powerpoint_app.planning.service import plan_with_provider
    from powerpoint_app.planning.providers import ProviderResult
    from powerpoint_app.domain.documents import Document, Block
    create_project(tmp_path)
    profile = pilot().teaching_profile
    save_profile(tmp_path, profile)
    documents = [Document('dc','sources/dc.md','DC',[Block('paragraph','24 V')])]
    class Provider:
        name = 'fake'; model = 'test'
        def __init__(self, answers): self.answers = iter(answers); self.calls = 0
        def complete(self, prompt):
            self.calls += 1
            return ProviderResult(next(self.answers), 10, 10)
    plain = load_plan(Path('examples/demo_project/slide-plan.json')).model_dump_json()
    provider = Provider([plain, pilot().model_dump_json()])
    plan, usage = plan_with_provider(tmp_path, documents, {'topic':'DC'}, provider)
    assert usage.calls == 2 and plan.teaching_profile == profile
    _, cached = plan_with_provider(tmp_path, documents, {'topic':'DC'}, Provider([]))
    assert cached.cache_hit and cached.calls == 0 and cached.input_tokens == 0 and cached.output_tokens == 0
    profile.required_method = 'En ændret metode'; save_profile(tmp_path, profile)
    updated = pilot(); updated.teaching_profile = profile
    next_provider = Provider([updated.model_dump_json()])
    _, usage = plan_with_provider(tmp_path, documents, {'topic':'DC'}, next_provider)
    assert next_provider.calls == 1 and not usage.cache_hit


def test_formula_links_to_real_branch_and_reveal_highlights_it(tmp_path):
    plan = pilot()
    bad = plan.model_dump(); bad['slides'][2]['elements'][1]['branch_index'] = 4
    with pytest.raises(ValueError, match='eksisterende gren'): SlidePlan.model_validate(bad)
    path = PptxRenderer(ROOT).render(plan, tmp_path/'linked.pptx')
    prs = Presentation(path)
    group = next(s for s in prs.slides[3].shapes if s.name == 'c3')
    from pptx.dml.color import RGBColor
    accent = RGBColor.from_string(PptxRenderer(ROOT).theme.accent)
    assert any(s.line.color.type is not None and s.line.color.rgb == accent for s in group.shapes if s.shape_type == 9)
