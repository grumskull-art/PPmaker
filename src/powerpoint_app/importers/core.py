from __future__ import annotations

import re
from pathlib import Path

from powerpoint_app.domain.documents import Block, Document


def import_document(path: Path, source_id: str) -> Document:
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        return import_markdown(path, source_id)
    if suffix == ".docx":
        return import_docx(path, source_id)
    if suffix == ".pdf":
        return import_pdf(path, source_id)
    raise ValueError(f"Ikke-understøttet filtype: {suffix}")


def import_markdown(path: Path, source_id: str) -> Document:
    text = path.read_text(encoding="utf-8-sig")
    blocks: list[Block] = []
    paragraph: list[str] = []
    lines = text.splitlines()

    def flush(line_no: int):
        if paragraph:
            blocks.append(Block("paragraph", "\n".join(paragraph).strip(), f"linje {line_no-len(paragraph)}–{line_no-1}"))
            paragraph.clear()

    for no, line in enumerate(lines, 1):
        if match := re.match(r"^(#{1,6})\s+(.+)$", line):
            flush(no)
            blocks.append(Block("heading", match.group(2).strip(), f"linje {no}"))
        elif line.strip():
            paragraph.append(line.rstrip())
        else:
            flush(no)
    flush(len(lines) + 1)
    title = next((b.text for b in blocks if b.kind == "heading"), path.stem)
    return Document(source_id, path.name, title, blocks)


def import_docx(path: Path, source_id: str) -> Document:
    from docx import Document as DocxDocument

    doc = DocxDocument(path)
    blocks: list[Block] = []
    warnings: list[str] = []
    for no, paragraph in enumerate(doc.paragraphs, 1):
        text = paragraph.text.strip()
        if text:
            kind = "heading" if paragraph.style and paragraph.style.name.startswith("Heading") else "paragraph"
            blocks.append(Block(kind, text, f"afsnit {no}"))
        xml = paragraph._p.xml
        if "m:oMath" in xml or "m:oMathPara" in xml:
            warnings.append(f"Avanceret Word-ligning ved afsnit {no} kræver manuel kontrol.")
    for no, table in enumerate(doc.tables, 1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        blocks.append(Block("table", locator=f"tabel {no}", rows=rows))
    title = next((b.text for b in blocks if b.kind == "heading"), path.stem)
    return Document(source_id, path.name, title, blocks, warnings)


def import_pdf(path: Path, source_id: str) -> Document:
    from pypdf import PdfReader

    reader = PdfReader(path)
    blocks: list[Block] = []
    warnings: list[str] = []
    extracted = 0
    for no, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        extracted += len(text)
        if text:
            blocks.append(Block("paragraph", text, f"side {no}"))
        else:
            warnings.append(f"Side {no} indeholder ingen udtrækkelig tekst; OCR kan være nødvendig.")
    if extracted < max(40, len(reader.pages) * 15):
        warnings.append("PDF'en ligner en billedscanning eller har meget lidt tekst. OCR er nødvendig før en dækkende slideplan.")
    return Document(source_id, path.name, path.stem, blocks, warnings)
